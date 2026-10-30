import os
import io
import csv
import json
import sqlite3
import secrets
from functools import wraps

from flask import (
    Flask, request, session, redirect, url_for,
    render_template, jsonify, Response, g,
)

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #

ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "hana-admin")
INGEST_TOKEN = os.environ.get("INGEST_TOKEN", "hana-secret-token")
SECRET_KEY = os.environ.get("SECRET_KEY", secrets.token_hex(32))
DB_PATH = os.environ.get("HANA_DB_PATH", os.path.join(os.path.dirname(__file__), "hana.db"))

app = Flask(__name__)
app.config["SECRET_KEY"] = SECRET_KEY
INTERACTION_FIELDS = [
    "client_ts", "question", "answer", "score", "stress_level",
    "Ea", "Ec", "Ad", "Be", "D", "Ep", "empathy_mode",
]

AEM_SUMMARY_FIELDS = [
    "language", "session_index", "checkin", "dass_raw", "S", "Ed", "H",
    "Ed_adaptive", "Bl", "pathway", "activity", "post_activity",
    # Revised single-session update: experimental condition, per-cycle
    # H / Ed_adaptive / Bl, checkpoint responses and the final check.
    "condition", "cycle2_choice",
    "H_1", "H_2", "H_3", "Ed_adaptive_1", "Ed_adaptive_2", "Ed_adaptive_3",
    "Bl_1", "Bl_2", "Bl_3", "checkpoint_1", "checkpoint_2", "final_check",
    "session_start", "participant_id",
]

# One row per interaction cycle (section 19 cycle-level logging).
AEM_CYCLE_FIELDS = ["cycle", "started_at", "Ed", "H", "Ed_adaptive", "Bl", "check_type", "check"]

# Session outcome fields kept in aem_history alongside the AEM values. The
# triadic build sends these under the legacy summary keys, so map
# history column -> payload key.
AEM_HISTORY_SESSION_FIELDS = [
    ("dass_score", "start_stress_score"),
    ("intention", "intention_level"),
    ("bl_path", "empathy_mode"),
    ("techniques", "techniques_used"),
    ("total_interactions", "total_interactions"),
    ("session_duration_minutes", "session_duration_minutes"),
    ("end_action", "end_action"),
    ("counselling_offered", "support_need_announced"),
]


def get_db():
    if "db" not in g:
        # timeout: wait up to 30s for a competing writer instead of hanging the
        # (single, on PythonAnywhere) worker forever — SQLite locking over the
        # host's network filesystem can stall under the client's retry storm.
        g.db = sqlite3.connect(DB_PATH, timeout=30)
        g.db.row_factory = sqlite3.Row
        # WAL lets the dashboard read while the game is writing, so reads don't
        # block writes (and vice-versa).
        g.db.execute("PRAGMA journal_mode=WAL")
        g.db.execute("PRAGMA busy_timeout=30000")
    return g.db


@app.teardown_appcontext
def close_db(exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = sqlite3.connect(DB_PATH, timeout=30)
    db.execute("PRAGMA journal_mode=WAL")
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS interactions (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            received_at     TEXT DEFAULT (datetime('now')),
            device_id       TEXT,
            app_type        TEXT,
            user_name       TEXT,
            session_id      TEXT,
            session_file    TEXT,
            client_ts       TEXT,
            question        TEXT,
            answer          TEXT,
            score           TEXT,
            stress_level    TEXT,
            Ea              REAL,
            Ec              REAL,
            Ad              REAL,
            Be              REAL,
            D               REAL,
            Ep              REAL,
            empathy_mode    TEXT,
            extra           TEXT
        );

        CREATE TABLE IF NOT EXISTS summaries (
            id                       INTEGER PRIMARY KEY AUTOINCREMENT,
            received_at              TEXT DEFAULT (datetime('now')),
            device_id                TEXT,
            app_type                 TEXT,
            user_name                TEXT,
            session_id               TEXT,
            session_file             TEXT,
            client_ts                TEXT,
            start_stress_score       TEXT,
            end_stress_score         TEXT,
            stress_improvement       TEXT,
            final_calm_rating        TEXT,
            modes_used               TEXT,
            techniques_used          TEXT,
            total_interactions       TEXT,
            session_duration_minutes TEXT,
            end_action               TEXT,
            stress_level             TEXT,
            empathy_mode             TEXT,
            intention_level          TEXT,
            empathy_E                TEXT,
            support_need_announced   TEXT,
            attempts                 INTEGER DEFAULT 1
        );

        CREATE INDEX IF NOT EXISTS idx_inter_device  ON interactions(device_id);
        CREATE INDEX IF NOT EXISTS idx_inter_session ON interactions(session_id);
        CREATE INDEX IF NOT EXISTS idx_summ_device   ON summaries(device_id);
        """
    )

    # Lightweight migration: CREATE TABLE IF NOT EXISTS won't add new columns to a
    # summaries table that already exists on a deployed DB, so add any that are
    # missing. New columns default to NULL for older rows.
    existing_cols = {row[1] for row in db.execute("PRAGMA table_info(summaries)")}
    for col in ("stress_level", "empathy_mode", "intention_level",
                "empathy_E", "support_need_announced", "app_type",
                "profile_slot", "profile_id"):
        if col not in existing_cols:
            db.execute("ALTER TABLE summaries ADD COLUMN %s TEXT" % col)
    if "attempts" not in existing_cols:
        db.execute("ALTER TABLE summaries ADD COLUMN attempts INTEGER DEFAULT 1")
    # AEM session values (HANA Update 01/10/26, step 17) for the latest session.
    for col in AEM_SUMMARY_FIELDS:
        if col not in existing_cols:
            db.execute("ALTER TABLE summaries ADD COLUMN %s TEXT" % col)

    # Interaction History Repository: one row per session, never overwritten,
    # so every session's S, Ed, H, Ed_adaptive and Bl is kept.
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS aem_history (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            received_at     TEXT DEFAULT (datetime('now')),
            device_id       TEXT,
            app_type        TEXT,
            user_name       TEXT,
            profile_id      TEXT,
            profile_slot    TEXT,
            session_id      TEXT,
            client_ts       TEXT,
            stress_level    TEXT,
            language        TEXT,
            session_index   TEXT,
            checkin         TEXT,
            dass_raw        TEXT,
            S               TEXT,
            Ed              TEXT,
            H               TEXT,
            Ed_adaptive     TEXT,
            Bl              TEXT,
            pathway         TEXT,
            activity        TEXT,
            post_activity   TEXT
        )
        """
    )
    hist_cols = {row[1] for row in db.execute("PRAGMA table_info(aem_history)")}
    for col in AEM_SUMMARY_FIELDS + [c for c, _ in AEM_HISTORY_SESSION_FIELDS]:
        if col not in hist_cols:
            db.execute("ALTER TABLE aem_history ADD COLUMN %s TEXT" % col)

    # Cycle-level history: one row per interaction cycle, never overwritten.
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS aem_cycles (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            received_at     TEXT DEFAULT (datetime('now')),
            device_id       TEXT,
            app_type        TEXT,
            condition       TEXT,
            user_name       TEXT,
            profile_id      TEXT,
            profile_slot    TEXT,
            session_id      TEXT,
            client_ts       TEXT,
            stress_level    TEXT,
            dass_raw        TEXT,
            S               TEXT,
            cycle           TEXT,
            started_at      TEXT,
            Ed              TEXT,
            H               TEXT,
            Ed_adaptive     TEXT,
            Bl              TEXT,
            check_type      TEXT,
            "check"         TEXT
        )
        """
    )
    cycle_cols = {row[1] for row in db.execute("PRAGMA table_info(aem_cycles)")}
    if "participant_id" not in cycle_cols:
        db.execute("ALTER TABLE aem_cycles ADD COLUMN participant_id TEXT")

    # Same for interactions: tag each row with the study arm + profile id/slot.
    inter_cols = {row[1] for row in db.execute("PRAGMA table_info(interactions)")}
    if "app_type" not in inter_cols:
        db.execute("ALTER TABLE interactions ADD COLUMN app_type TEXT")
    if "profile_slot" not in inter_cols:
        db.execute("ALTER TABLE interactions ADD COLUMN profile_slot TEXT")
    if "profile_id" not in inter_cols:
        db.execute("ALTER TABLE interactions ADD COLUMN profile_id TEXT")

    db.commit()
    db.close()


# --------------------------------------------------------------------------- #
# CORS — the itch.io web build runs the game from a different origin
# (e.g. https://html.itch.zone), so the browser will only let it POST to the
# /api/ endpoints if we answer the preflight and echo the allow headers. Auth is
# by the X-Ingest-Token header (no cookies), so a wildcard origin is safe here.
# Override with the HANA_CORS_ORIGIN env var to lock it to a specific origin.
# --------------------------------------------------------------------------- #

CORS_ORIGIN = os.environ.get("HANA_CORS_ORIGIN", "*")


@app.before_request
def _cors_preflight():
    if request.method == "OPTIONS" and request.path.startswith("/api/"):
        return ("", 204)


@app.after_request
def _cors_headers(resp):
    if request.path.startswith("/api/"):
        resp.headers["Access-Control-Allow-Origin"] = CORS_ORIGIN
        resp.headers["Access-Control-Allow-Headers"] = "Content-Type, X-Ingest-Token"
        resp.headers["Access-Control-Allow-Methods"] = "POST, OPTIONS"
        resp.headers["Access-Control-Max-Age"] = "86400"
    return resp


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("admin"):
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)
    return wrapped


def check_ingest_token(req):
    token = req.headers.get("X-Ingest-Token", "")
    return secrets.compare_digest(token, INGEST_TOKEN)


@app.route("/api/interaction", methods=["POST"])
def api_interaction():
    if not check_ingest_token(request):
        return jsonify(ok=False, error="bad token"), 401

    data = request.get_json(silent=True) or {}

    # Anything outside the known columns is stashed in `extra` so we never lose
    # the enhanced-logging fields (Ea_history, struggle metrics, etc.).
    known = {"device_id", "app_type", "user_name", "profile_id", "profile_slot", "session_id", "session_file"} | set(INTERACTION_FIELDS)
    extra = {k: v for k, v in data.items() if k not in known}

    db = get_db()
    db.execute(
        """
        INSERT INTO interactions
            (device_id, app_type, user_name, profile_id, profile_slot, session_id, session_file, client_ts,
             question, answer, score, stress_level,
             Ea, Ec, Ad, Be, D, Ep, empathy_mode, extra)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            data.get("device_id"), data.get("app_type"),
            data.get("user_name"), _s(data.get("profile_id")), _s(data.get("profile_slot")), data.get("session_id"),
            data.get("session_file"), data.get("client_ts"),
            data.get("question"), data.get("answer"), _s(data.get("score")),
            data.get("stress_level"),
            _f(data.get("Ea")), _f(data.get("Ec")), _f(data.get("Ad")),
            _f(data.get("Be")), _f(data.get("D")), _f(data.get("Ep")),
            data.get("empathy_mode"),
            json.dumps(extra) if extra else None,
        ),
    )
    db.commit()
    return jsonify(ok=True)


@app.route("/api/summary", methods=["POST"])
def api_summary():
    if not check_ingest_token(request):
        return jsonify(ok=False, error="bad token"), 401

    data = request.get_json(silent=True) or {}
    db = get_db()

    # Each profile gets a stable uuid (profile_id) when it's created in a slot;
    # that's the per-profile identity we dedup on. profile_slot is kept too, but
    # only as informational context (it gets reused when a slot is reset).
    profile_id = data.get("profile_id")
    profile_id = str(profile_id) if profile_id is not None and str(profile_id).strip() != "" else None
    slot = data.get("profile_slot")
    slot = str(slot) if slot is not None and str(slot).strip() != "" else None

    fields = (
        profile_id, slot,
        data.get("session_id"), data.get("session_file"), data.get("client_ts"),
        _s(data.get("start_stress_score")), _s(data.get("end_stress_score")),
        _s(data.get("stress_improvement")), _s(data.get("final_calm_rating")),
        data.get("modes_used"), data.get("techniques_used"),
        _s(data.get("total_interactions")), _s(data.get("session_duration_minutes")),
        data.get("end_action"),
        data.get("stress_level"), data.get("empathy_mode"),
        data.get("intention_level"), _s(data.get("empathy_E")),
        _s(data.get("support_need_announced")),
    )

    # One row per profile: replaying the SAME profile overwrites its summary
    # (attempts++), while a DIFFERENT profile — even same username, or a new
    # profile created after a slot was reset — gets a fresh profile_id and its
    # own row, so reset+reuse never clobbers earlier data. Fall back to session_id
    # for builds that don't send a profile_id yet (those never overwrite either).
    if profile_id is not None:
        existing = db.execute(
            "SELECT id, attempts FROM summaries WHERE device_id = ? AND profile_id = ?",
            (data.get("device_id"), profile_id),
        ).fetchone()
    else:
        existing = db.execute(
            "SELECT id, attempts FROM summaries WHERE device_id = ? AND session_id = ?",
            (data.get("device_id"), data.get("session_id")),
        ).fetchone()

    if existing:
        db.execute(
            """
            UPDATE summaries SET
                received_at = datetime('now'),
                app_type = ?,
                profile_id = ?, profile_slot = ?,
                session_id = ?, session_file = ?, client_ts = ?,
                start_stress_score = ?, end_stress_score = ?, stress_improvement = ?,
                final_calm_rating = ?, modes_used = ?, techniques_used = ?,
                total_interactions = ?, session_duration_minutes = ?, end_action = ?,
                stress_level = ?, empathy_mode = ?, intention_level = ?, empathy_E = ?,
                support_need_announced = ?,
                attempts = ?
            WHERE id = ?
            """,
            (data.get("app_type"),) + fields + ((existing["attempts"] or 1) + 1, existing["id"]),
        )
    else:
        db.execute(
            """
            INSERT INTO summaries
                (device_id, app_type, user_name, profile_id, profile_slot, session_id, session_file,
                 client_ts, start_stress_score, end_stress_score, stress_improvement,
                 final_calm_rating, modes_used, techniques_used,
                 total_interactions, session_duration_minutes, end_action,
                 stress_level, empathy_mode, intention_level, empathy_E,
                 support_need_announced, attempts)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,1)
            """,
            (data.get("device_id"), data.get("app_type"), data.get("user_name")) + fields,
        )

    if data.get("Bl") is not None or data.get("Bl_1") is not None:
        aem_values = tuple(_s(data.get(k)) for k in AEM_SUMMARY_FIELDS)
        row_id = existing["id"] if existing else db.execute("SELECT last_insert_rowid()").fetchone()[0]
        db.execute(
            "UPDATE summaries SET %s WHERE id = ?" % ", ".join("%s = ?" % k for k in AEM_SUMMARY_FIELDS),
            aem_values + (row_id,),
        )
        hist_cols = AEM_SUMMARY_FIELDS + [c for c, _ in AEM_HISTORY_SESSION_FIELDS]
        db.execute(
            """
            INSERT INTO aem_history
                (device_id, app_type, user_name, profile_id, profile_slot, session_id,
                 client_ts, stress_level, %s)
            VALUES (?,?,?,?,?,?,?,?,%s)
            """ % (", ".join(hist_cols), ",".join("?" * len(hist_cols))),
            (data.get("device_id"), data.get("app_type"), data.get("user_name"),
             profile_id, slot, data.get("session_id"), data.get("client_ts"),
             data.get("stress_level")) + aem_values
            + tuple(_s(data.get(k)) for _, k in AEM_HISTORY_SESSION_FIELDS),
        )
    cycles = data.get("cycles")
    if isinstance(cycles, list):
        for c in cycles:
            if not isinstance(c, dict):
                continue
            db.execute(
                """
                INSERT INTO aem_cycles
                    (device_id, app_type, condition, user_name, profile_id, profile_slot,
                     participant_id, session_id, client_ts, stress_level, dass_raw, S,
                     cycle, started_at, Ed, H, Ed_adaptive, Bl, check_type, "check")
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (data.get("device_id"), data.get("app_type"), data.get("condition"),
                 data.get("user_name"), profile_id, slot, _s(data.get("participant_id")),
                 data.get("session_id"),
                 data.get("client_ts"), data.get("stress_level"),
                 _s(data.get("dass_raw")), _s(data.get("S")))
                + tuple(_s(c.get(k)) for k in AEM_CYCLE_FIELDS),
            )
    db.commit()
    return jsonify(ok=True)


@app.route("/api/health")
def health():
    return jsonify(ok=True)


def _f(v):
    """Best-effort float; keeps NULLs as NULL."""
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _s(v):
    """Store scalars as text without turning None into the string 'None'."""
    if v is None:
        return None
    return str(v)


@app.template_filter("num")
def _num(v):
    """Show model values to 3 decimals on the admin pages; blanks stay blank."""
    if v is None or v == "":
        return ""
    if isinstance(v, float):
        return "%.3f" % v
    return v

@app.route("/")
def index():
    return redirect(url_for("dashboard"))


@app.route("/login", methods=["GET", "POST"])
def login():
    error = None
    if request.method == "POST":
        if secrets.compare_digest(request.form.get("password", ""), ADMIN_PASSWORD):
            session["admin"] = True
            return redirect(request.args.get("next") or url_for("dashboard"))
        error = "Wrong password."
    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/dashboard")
@login_required
def dashboard():
    db = get_db()
    users = db.execute(
        """
        SELECT
            device_id,
            MAX(user_name)                          AS user_name,
            COUNT(*)                                AS interactions,
            COUNT(DISTINCT session_id)              AS sessions,
            MAX(received_at)                        AS last_seen
        FROM interactions
        GROUP BY device_id
        ORDER BY last_seen DESC
        """
    ).fetchall()

    totals = db.execute(
        "SELECT COUNT(*) AS i, COUNT(DISTINCT device_id) AS d, "
        "COUNT(DISTINCT session_id) AS s FROM interactions"
    ).fetchone()
    summ_total = db.execute("SELECT COUNT(*) AS c FROM summaries").fetchone()["c"]

    return render_template(
        "dashboard.html", users=users, totals=totals, summ_total=summ_total
    )


@app.route("/user/<device_id>")
@login_required
def user_detail(device_id):
    db = get_db()
    sessions = db.execute(
        """
        SELECT
            session_id,
            MAX(user_name)  AS user_name,
            COUNT(*)        AS turns,
            MIN(client_ts)  AS started,
            MAX(client_ts)  AS ended
        FROM interactions
        WHERE device_id = ?
        GROUP BY session_id
        ORDER BY started DESC
        """,
        (device_id,),
    ).fetchall()

    summaries = db.execute(
        "SELECT * FROM summaries WHERE device_id = ? ORDER BY received_at DESC",
        (device_id,),
    ).fetchall()

    name = sessions[0]["user_name"] if sessions else device_id
    return render_template(
        "user.html", device_id=device_id, name=name,
        sessions=sessions, summaries=summaries,
    )


# --------------------------------------------------------------------------- #
# Per-session log in the same one-file layout the game writes on the device
# (game/script.rpy SESSION_LOG_HEADER / SESSION_SUMMARY_FIELDS), so a session
# downloaded here can be compared line for line with the tester's local file.
# --------------------------------------------------------------------------- #

SESSION_LOG_HEADER = [
    "Participant_ID", "Session_ID", "Condition",
    "Timestamp", "Cycle", "Stage", "Duration",
    "HANA_Question", "User_Response", "HANA_Response",
    "Ea", "Ec", "Be", "D", "Ep", "E",
    "I",
    "Ed", "H", "Ed_Adaptive",
    "Behaviour_Profile",
]

# Study arm codes, as in the game's log file name and Condition column.
CONDITION_CODES = {"adaptive": "ADP", "strict_empathy": "TRI", "neutral": "NEU"}

# (label in the game's file, aem_history column)
SESSION_SUMMARY_FIELDS = [
    ("Session start", "session_start"), ("Session end", "client_ts"),
    ("User", "user_name"), ("Participant ID", "participant_id"), ("Session ID", "session_id"),
    ("Condition", "condition"),
    ("Check-in", "checkin"), ("DASS stress raw (0-21)", "dass_raw"),
    ("DASS stress score (0-42)", "dass_score"), ("Stress level", "stress_level"),
    ("S", "S"), ("Ed", "Ed"), ("H", "H"), ("Ed_adaptive", "Ed_adaptive"), ("Bl", "Bl"),
    ("Intention (I)", "intention"), ("Pathway", "pathway"),
    ("Cycle 2 choice", "cycle2_choice"), ("Activities", "activity"),
    ("Checkpoint 1", "checkpoint_1"), ("Checkpoint 2", "checkpoint_2"),
    ("Final check", "final_check"), ("Bl per cycle", "bl_path"),
    ("Counselling offered", "counselling_offered"), ("Turns", "total_interactions"),
]


def _session_log_rows(db, device_id, session_id):
    """Turn rows then summary rows, as dicts keyed by SESSION_LOG_HEADER, plus
    the user's name and the log's file name on the device."""
    turns = []
    for r in db.execute(
        "SELECT * FROM interactions WHERE device_id = ? AND session_id = ? "
        "ORDER BY client_ts, id",
        (device_id, session_id),
    ):
        r = dict(r)
        try:
            x = json.loads(r.get("extra") or "{}")
        except (TypeError, ValueError):
            x = {}
        if not isinstance(x, dict):
            x = {}
        turns.append((r, x))
    summary = db.execute(
        "SELECT * FROM aem_history WHERE device_id = ? AND session_id = ? "
        "ORDER BY id DESC LIMIT 1",
        (device_id, session_id),
    ).fetchone()
    summary = dict(summary) if summary else {}

    rows = [_log_row(r, x) for r, x in turns]
    if summary:
        participant_id = summary.get("participant_id")
        code = CONDITION_CODES.get(summary.get("condition"), summary.get("condition"))
        for label, col in SESSION_SUMMARY_FIELDS:
            value = code if col == "condition" else summary.get(col)
            rows.append({
                "Participant_ID": participant_id, "Session_ID": summary["session_id"],
                "Condition": code, "Timestamp": summary["client_ts"], "Cycle": "",
                "Stage": "Summary", "Duration": "",
                "HANA_Question": label, "User_Response": value,
            })
    name = turns[0][0]["user_name"] if turns else summary.get("user_name", "")
    session_file = turns[0][0]["session_file"] if turns else ""
    return rows, name, session_file


def _log_row(r, x):
    """One interactions row (plus its `extra` JSON) in SESSION_LOG_HEADER layout."""
    return {
        "Participant_ID": x.get("participant_id"), "Session_ID": r["session_id"],
        "Condition": CONDITION_CODES.get(x.get("condition"), x.get("condition")),
        "Timestamp": r["client_ts"], "Cycle": x.get("cycle"), "Stage": x.get("stage"),
        "Duration": x.get("duration"),
        "HANA_Question": r["question"], "User_Response": r["answer"],
        "HANA_Response": x.get("hana_response"),
        "Ea": r["Ea"], "Ec": r["Ec"], "Be": r["Be"], "D": r["D"], "Ep": r["Ep"],
        "E": x.get("E"), "I": x.get("I"),
        "Ed": x.get("Ed"), "H": x.get("H"), "Ed_Adaptive": x.get("Ed_adaptive"),
        "Behaviour_Profile": x.get("Bl"),
    }


@app.route("/session/<device_id>/<session_id>")
@login_required
def session_detail(device_id, session_id):
    rows, name, session_file = _session_log_rows(get_db(), device_id, session_id)
    return render_template(
        "session.html", rows=rows, device_id=device_id,
        session_id=session_id, name=name, session_file=session_file,
    )


@app.route("/session/<device_id>/<session_id>.csv")
@login_required
def session_log_csv(device_id, session_id):
    rows, name, session_file = _session_log_rows(get_db(), device_id, session_id)
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(SESSION_LOG_HEADER)
    for row in rows:
        writer.writerow([row.get(k) for k in SESSION_LOG_HEADER])
    # Same file name as on the device, e.g. HANA_P01_ADP_20261026_144717.csv
    filename = session_file or "session_%s.csv" % session_id
    return Response(
        out.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=%s" % filename},
    )


# The study arms, as stored in `app_type`, with display labels named as in the
# study design (Adaptive / Triadic / Baseline) and the log file codes. The
# stored values are older names: `triadic` is the Adaptive arm, and
# `strict_empathy` the Triadic one. `adaptive` is the separate older HANAAdaptive app.
ARM_LABELS = [
    ("triadic", "Adaptive (ADP)"),
    ("strict_empathy", "Triadic (TRI)"),
    ("neutral", "Baseline (NEU)"),
    ("adaptive", "Adaptive, old app"),
]
ARM_SHORT = {"triadic": "Adaptive", "strict_empathy": "Triadic", "neutral": "Baseline", "adaptive": "Adaptive old app"}


@app.route("/summaries")
@login_required
def summaries():
    db = get_db()

    # Optional ?arm=<adaptive|strict_empathy|neutral> filter — "all" (or missing)
    # shows everything.
    arm = request.args.get("arm", "all")
    valid_arms = {k for k, _ in ARM_LABELS}
    # Order by the device timestamp (when the session actually happened), falling
    # back to server-receive time for rows that predate client_ts. client_ts is
    # "%Y-%m-%d %H:%M:%S" text, so it sorts chronologically as a string.
    if arm in valid_arms:
        rows = db.execute(
            "SELECT * FROM summaries WHERE app_type = ? "
            "ORDER BY COALESCE(client_ts, received_at) DESC",
            (arm,),
        ).fetchall()
    else:
        arm = "all"
        rows = db.execute(
            "SELECT * FROM summaries "
            "ORDER BY COALESCE(client_ts, received_at) DESC"
        ).fetchall()

    # Per-arm counts for the filter tabs / stat cards.
    counts = {k: 0 for k, _ in ARM_LABELS}
    total = 0
    for r in db.execute("SELECT app_type, COUNT(*) AS c FROM summaries GROUP BY app_type"):
        total += r["c"]
        if r["app_type"] in counts:
            counts[r["app_type"]] = r["c"]

    return render_template(
        "summaries.html", rows=rows, arm=arm,
        arm_labels=ARM_LABELS, arm_short=ARM_SHORT, counts=counts, total=total,
    )


@app.route("/summary/<int:summary_id>/delete", methods=["POST"])
@login_required
def delete_summary(summary_id):
    db = get_db()
    db.execute("DELETE FROM summaries WHERE id = ?", (summary_id,))
    db.commit()
    return redirect(request.referrer or url_for("summaries"))


@app.route("/summaries/delete", methods=["POST"])
@login_required
def delete_summaries_bulk():
    """Delete every summary in the current view: one arm, or all of them."""
    arm = request.form.get("arm", "all")
    db = get_db()
    if arm in {k for k, _ in ARM_LABELS}:
        db.execute("DELETE FROM summaries WHERE app_type = ?", (arm,))
        db.commit()
        return redirect(url_for("summaries", arm=arm))
    db.execute("DELETE FROM summaries")
    db.commit()
    return redirect(url_for("summaries"))


@app.route("/user/<device_id>/delete", methods=["POST"])
@login_required
def delete_user(device_id):
    """Remove a device/user entirely: all their interactions, summaries and
    AEM history, so they also drop out of the Excel exports."""
    db = get_db()
    for table in ("interactions", "summaries", "aem_history", "aem_cycles"):
        db.execute("DELETE FROM %s WHERE device_id = ?" % table, (device_id,))
    db.commit()
    return redirect(url_for("dashboard"))


@app.route("/export/interactions.csv")
@login_required
def export_interactions():
    return _export("interactions")


@app.route("/export/summaries.csv")
@login_required
def export_summaries():
    return _export("summaries")


@app.route("/export/aem_history.csv")
@login_required
def export_aem_history():
    return _export("aem_history")


@app.route("/export/aem_cycles.csv")
@login_required
def export_aem_cycles():
    return _export("aem_cycles")


# --------------------------------------------------------------------------- #
# Excel exports. Each sheet lists only the columns that arm actually fills,
# with readable headers, as (db column, header) pairs.
# --------------------------------------------------------------------------- #

ARM_SHEET_ORDER = [k for k, _ in ARM_LABELS]

# Adaptive Triadic: one row per session from aem_history (never overwritten,
# unlike `summaries`, which keeps only the latest session per profile).
TRIADIC_SESSION_COLUMNS = [
    ("received_at", "Received (server, UTC)"),
    ("session_start", "Session start (device)"),
    ("client_ts", "Session end (device)"),
    ("participant_id", "Participant ID"),
    ("user_name", "User"),
    ("profile_id", "Profile ID"),
    ("device_id", "Device ID"),
    ("session_id", "Session ID"),
    ("session_index", "Session #"),
    ("condition", "Condition"),
    ("language", "Language"),
    ("checkin", "Check-in"),
    ("dass_raw", "DASS stress raw (0-21)"),
    ("dass_score", "DASS stress score (0-42)"),
    ("stress_level", "Stress level"),
    ("S", "S"),
    ("Ed", "Ed"),
    ("H", "H"),
    ("Ed_adaptive", "Ed_adaptive"),
    ("Bl", "Bl"),
    ("intention", "Intention (I)"),
    ("pathway", "Pathway"),
    ("cycle2_choice", "Cycle 2 choice"),
    ("activity", "Activities"),
    ("checkpoint_1", "Checkpoint 1"),
    ("checkpoint_2", "Checkpoint 2"),
    ("final_check", "Final check"),
    ("bl_path", "Bl per cycle"),
    ("counselling_offered", "Counselling offered"),
    ("total_interactions", "Turns"),
    ("session_duration_minutes", "Duration (min)"),
    ("end_action", "Ended by"),
]

TRIADIC_CYCLE_COLUMNS = [
    ("received_at", "Received (server, UTC)"),
    ("client_ts", "Session time (device)"),
    ("participant_id", "Participant ID"),
    ("user_name", "User"),
    ("profile_id", "Profile ID"),
    ("device_id", "Device ID"),
    ("session_id", "Session ID"),
    ("condition", "Condition"),
    ("stress_level", "Stress level"),
    ("dass_raw", "DASS stress raw (0-21)"),
    ("S", "S"),
    ("cycle", "Cycle"),
    ("started_at", "Cycle started (device)"),
    ("Ed", "Ed"),
    ("H", "H"),
    ("Ed_adaptive", "Ed_adaptive"),
    ("Bl", "Bl"),
    ("check_type", "Check"),
    ("check", "Answer"),
]

# The older arms (adaptive / strict_empathy / neutral) use the original
# summary layout; the AEM columns stay empty for them, so they're left out.
LEGACY_SUMMARY_COLUMNS = [
    ("received_at", "Received (server, UTC)"),
    ("client_ts", "Session time (device)"),
    ("user_name", "User"),
    ("profile_id", "Profile ID"),
    ("device_id", "Device ID"),
    ("session_id", "Session ID"),
    ("attempts", "Attempts"),
    ("start_stress_score", "Start stress"),
    ("end_stress_score", "End stress"),
    ("stress_improvement", "Improvement"),
    ("final_calm_rating", "Calm rating"),
    ("modes_used", "Modes used"),
    ("techniques_used", "Techniques"),
    ("total_interactions", "Turns"),
    ("session_duration_minutes", "Duration (min)"),
    ("end_action", "Ended by"),
    ("stress_level", "Stress level"),
    ("empathy_mode", "Empathy mode"),
    ("intention_level", "Intention"),
    ("empathy_E", "E"),
    ("support_need_announced", "Support announced"),
]

# Combined sheet: every session from every arm. Shared measures come first,
# then the columns only one kind of arm fills.
ALL_ARMS_COLUMNS = [
    ("arm", "Arm"),
    ("session_start", "Session start (device)"),
    ("client_ts", "Session end (device)"),
    ("received_at", "Received (server, UTC)"),
    ("participant_id", "Participant ID"),
    ("user_name", "User"),
    ("profile_id", "Profile ID"),
    ("device_id", "Device ID"),
    ("session_id", "Session ID"),
    ("stress_level", "Stress level"),
    ("start_stress", "Start stress (DASS 0-42)"),
    ("end_stress_score", "End stress"),
    ("stress_improvement", "Improvement"),
    ("techniques", "Techniques"),
    ("intention", "Intention"),
    ("total_interactions", "Turns"),
    ("session_duration_minutes", "Duration (min)"),
    ("end_action", "Ended by"),
    # Adaptive Triadic only
    ("session_index", "Session #"),
    ("checkin", "Check-in"),
    ("dass_raw", "DASS stress raw (0-21)"),
    ("S", "S"),
    ("Ed", "Ed"),
    ("H", "H"),
    ("Ed_adaptive", "Ed_adaptive"),
    ("Bl", "Bl"),
    ("bl_path", "Bl per cycle"),
    ("pathway", "Pathway"),
    ("cycle2_choice", "Cycle 2 choice"),
    ("activity", "Activities"),
    ("checkpoint_1", "Checkpoint 1"),
    ("checkpoint_2", "Checkpoint 2"),
    ("final_check", "Final check"),
    ("counselling_offered", "Counselling offered"),
    # Older arms only
    ("final_calm_rating", "Calm rating"),
    ("modes_used", "Modes used"),
    ("empathy_mode", "Empathy mode"),
    ("empathy_E", "E"),
    ("support_need_announced", "Support announced"),
]

# Participant numbers restart on every device, so Device_ID comes first.
INTERACTION_COLUMNS = [("Device_ID", "Device_ID")] + [(k, k) for k in SESSION_LOG_HEADER] + [
    ("User_Name", "User_Name"), ("Received", "Received (server, UTC)"),
]

# Layout for rows from older builds, which don't send the log columns above.
OLDER_INTERACTION_COLUMNS = [
    ("received_at", "Received (server, UTC)"),
    ("client_ts", "Time (device)"),
    ("user_name", "User"),
    ("profile_id", "Profile ID"),
    ("device_id", "Device ID"),
    ("session_id", "Session ID"),
    ("question", "Question / event"),
    ("answer", "Answer"),
    ("score", "Score"),
    ("stress_level", "Stress level"),
    ("Ea", "Ea"), ("Ec", "Ec"), ("Ad", "Ad"), ("Be", "Be"), ("D", "D"), ("Ep", "Ep"),
    ("empathy_mode", "Mode"),
]
# Keys from the interactions `extra` JSON that get their own column first.
OLDER_INTERACTION_EXTRA_ORDER = [
    "cycle", "S", "E", "I", "Ed", "H", "Ed_adaptive", "Bl", "hana_response",
    "condition", "language",
]

# Stored as TEXT in SQLite; written to Excel as numbers so they can be charted.
NUMERIC_COLUMNS = {
    "session_index", "dass_raw", "dass_score", "start_stress", "S", "Ed", "H", "Ed_adaptive",
    "intention", "total_interactions", "session_duration_minutes", "cycle",
    "attempts", "start_stress_score", "end_stress_score", "stress_improvement",
    "empathy_E", "score", "Ea", "Ec", "Ad", "Be", "D", "Ep", "E", "I",
    "Cycle", "Duration", "Ed_Adaptive",
}


def _sheet_name(title):
    # Excel caps sheet names at 31 chars and forbids []:*?/\
    return "".join(ch for ch in title if ch not in "[]:*?/\\")[:31]


def _arm_title(arm):
    return _sheet_name(dict(ARM_LABELS).get(arm, arm or "untagged"))


def _filled_columns(columns, rows):
    """Drop columns no row fills; keep them all while a sheet has no rows."""
    if not rows:
        return columns
    return [(c, h) for c, h in columns if any(r.get(c) not in (None, "") for r in rows)]


def _cell(col, value):
    if col in NUMERIC_COLUMNS and value not in (None, ""):
        try:
            f = float(value)
            return int(f) if f.is_integer() else f
        except (TypeError, ValueError):
            pass
    if isinstance(value, (dict, list)):
        return json.dumps(value)
    return value


def _add_sheet(wb, title, columns, rows):
    from openpyxl.styles import Font

    ws = wb.create_sheet(title=title)
    ws.append([header for _, header in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([_cell(col, r.get(col)) for col, _ in columns])
    ws.freeze_panes = "A2"
    for i, (_, header) in enumerate(columns, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(10, min(40, len(header) + 2))
    return ws


def _rows_by_arm(rows):
    """Group rows by app_type, in ARM_SHEET_ORDER, with unknown arms last."""
    grouped = {}
    for r in rows:
        grouped.setdefault(r.get("app_type") or "untagged", []).append(r)
    order = [a for a in ARM_SHEET_ORDER if a in grouped]
    order += sorted(a for a in grouped if a not in order)
    return [(a, grouped[a]) for a in order]


def _xlsx_response(wb, filename):
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return Response(
        bio.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=%s" % filename},
    )


@app.route("/export/summaries.xlsx")
@login_required
def export_summaries_xlsx():
    from openpyxl import Workbook

    db = get_db()
    history = [dict(r) for r in db.execute(
        "SELECT * FROM aem_history ORDER BY COALESCE(client_ts, received_at), id")]
    cycles = [dict(r) for r in db.execute(
        "SELECT * FROM aem_cycles ORDER BY COALESCE(client_ts, received_at), id, cycle")]
    # Sessions from the current game (any arm) are in aem_history; the legacy
    # sheets keep only older builds' summaries, so nothing is counted twice.
    in_history = {(r["device_id"], r["session_id"]) for r in history}
    legacy = [dict(r) for r in db.execute(
        "SELECT * FROM summaries WHERE COALESCE(app_type, '') != 'triadic' "
        "ORDER BY COALESCE(client_ts, received_at), id")
        if (r["device_id"], r["session_id"]) not in in_history]

    # Bring both layouts onto the shared column names for the combined sheet.
    combined = [dict(r, arm=_arm_title(r.get("app_type") or "triadic"),
                     start_stress=r.get("dass_score")) for r in history]
    combined += [dict(r, arm=_arm_title(r.get("app_type")),
                      start_stress=r.get("start_stress_score"),
                      techniques=r.get("techniques_used"),
                      intention=r.get("intention_level")) for r in legacy]
    combined.sort(key=lambda r: r.get("client_ts") or r.get("received_at") or "")

    wb = Workbook()
    wb.remove(wb.active)
    _add_sheet(wb, "All arms", ALL_ARMS_COLUMNS, combined)
    # Sessions + cycles sheets per arm. The Adaptive Triadic pair always comes
    # first, even when still empty; the other arms only skip the AEM values
    # they never compute (strict empathy: Ed / H / Ed_adaptive, baseline: all).
    hist_by_arm = dict(_rows_by_arm([dict(r, app_type=r.get("app_type") or "triadic") for r in history]))
    cyc_by_arm = dict(_rows_by_arm([dict(r, app_type=r.get("app_type") or "triadic") for r in cycles]))
    for arm in ["triadic"] + [a for a in hist_by_arm if a != "triadic"]:
        short = ARM_SHORT.get(arm, arm)
        arm_hist, arm_cyc = hist_by_arm.get(arm, []), cyc_by_arm.get(arm, [])
        session_cols, cycle_cols = TRIADIC_SESSION_COLUMNS, TRIADIC_CYCLE_COLUMNS
        if arm != "triadic":
            session_cols = _filled_columns(session_cols, arm_hist)
            cycle_cols = _filled_columns(cycle_cols, arm_cyc)
        _add_sheet(wb, _sheet_name(short + " - sessions"), session_cols, arm_hist)
        _add_sheet(wb, _sheet_name(short + " - cycles"), cycle_cols, arm_cyc)
    for arm, rows in _rows_by_arm(legacy):
        _add_sheet(wb, _sheet_name(ARM_SHORT.get(arm, arm or "untagged") + " - older build"),
                   LEGACY_SUMMARY_COLUMNS, rows)
    return _xlsx_response(wb, "hana_summaries.xlsx")


@app.route("/export/interactions.xlsx")
@login_required
def export_interactions_xlsx():
    from openpyxl import Workbook

    db = get_db()
    rows, older = [], []
    for r in db.execute("SELECT * FROM interactions ORDER BY id ASC"):
        r = dict(r)
        try:
            extra = json.loads(r.get("extra") or "{}")
        except (TypeError, ValueError):
            extra = {}
        if not isinstance(extra, dict):
            extra = {}
        if "stage" in extra:
            # Current game: the same columns as its session log file.
            row = _log_row(r, extra)
            row.update(Device_ID=r["device_id"], User_Name=r["user_name"],
                       Received=r["received_at"], app_type=r["app_type"])
            rows.append(row)
        else:
            for k, v in extra.items():
                r.setdefault(k, v)
            older.append(r)

    wb = Workbook()
    wb.remove(wb.active)
    for arm, arm_rows in _rows_by_arm(rows):
        _add_sheet(wb, _arm_title(arm), INTERACTION_COLUMNS, arm_rows)
    # Rows from older builds (or logged before the current log layout) keep
    # their original columns, on their own sheet per arm.
    known = {c for c, _ in OLDER_INTERACTION_COLUMNS} | {
        "id", "app_type", "profile_slot", "session_file", "extra"}
    for arm, arm_rows in _rows_by_arm(older):
        # Each arm logs a different set of extra fields, so the columns are
        # built from what that arm's rows actually contain.
        extra_keys = {k for r in arm_rows for k in r if k not in known}
        ordered = [k for k in OLDER_INTERACTION_EXTRA_ORDER if k in extra_keys]
        ordered += sorted(extra_keys - set(ordered))
        _add_sheet(wb, _sheet_name(ARM_SHORT.get(arm, arm or "untagged") + " - older build"),
                   OLDER_INTERACTION_COLUMNS + [(k, k) for k in ordered], arm_rows)
    if not wb.sheetnames:
        _add_sheet(wb, "No data yet", INTERACTION_COLUMNS, [])
    return _xlsx_response(wb, "hana_interactions.xlsx")


def _export(table):
    db = get_db()
    rows = db.execute("SELECT * FROM %s ORDER BY id ASC" % table).fetchall()
    out = io.StringIO()
    writer = csv.writer(out)
    if rows:
        writer.writerow(rows[0].keys())
        for r in rows:
            writer.writerow([r[k] for k in r.keys()])
    return Response(
        out.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=%s.csv" % table},
    )


# Initialise the schema as soon as the module is imported, so it works both
# under `flask run`, `python app.py`, and WSGI hosting (gunicorn / PythonAnywhere).
init_db()


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=True)
